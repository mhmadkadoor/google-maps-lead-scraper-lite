"""
Local Lead Finder (Lite Edition) - Excel & CSV Exporters
"""

import os
from typing import List
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from models import Lead


def export_to_excel(leads: List[Lead], output_path: str) -> str:
    df = Lead.to_dataframe(leads)
    df.rename(columns={
        "name": "Business Name",
        "phone": "Phone Number",
        "email": "Email Address",
        "website": "Website URL",
        "rating": "Google Rating",
        "review_count": "Review Count",
        "address": "Full Address",
        "category": "Category",
        "google_maps_url": "Google Maps Link",
    }, inplace=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Leads")
        ws = writer.sheets["Leads"]

        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_idx in range(1, len(df.columns) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        ws.row_dimensions[1].height = 28
        ws.freeze_panes = "A2"

        for col_idx, col in enumerate(df.columns, 1):
            max_len = max(df[col].astype(str).map(len).max() if not df.empty else 0, len(str(col)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 4, 12), 40)

    return output_path


def export_to_csv(leads: List[Lead], output_path: str) -> str:
    df = Lead.to_dataframe(leads)
    df.to_csv(output_path, index=False, encoding="utf-8-sig")
    return output_path
